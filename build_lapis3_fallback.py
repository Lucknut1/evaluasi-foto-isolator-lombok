from __future__ import annotations

import csv
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "lapis3_config.json"


def read_json(path: Path, fallback=None):
    if fallback is None:
        fallback = []
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return fallback


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def score_band(row, config):
    value = row.get("relationship_score")
    if value in (None, ""):
        return "UNKNOWN"
    value = float(value)
    confidence = config["relationship"]["confidence"]
    if value >= confidence["very_high"]:
        return "VERY_HIGH"
    if value >= confidence["high"]:
        return "HIGH"
    if value >= confidence["medium"]:
        return "MEDIUM"
    return "LOW"


def setup_sheet(ws, title, headers, rows, widths=None):
    ws.sheet_view.showGridLines = False
    ws["A2"] = title
    ws["A2"].font = Font(name="Arial", size=15, bold=True, color="173B58")
    header_row = 4
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B58")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_no, row in enumerate(rows, header_row + 1):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row_no, col, row.get(header, ""))
            cell.font = Font(name="Arial", size=9, color="203042")
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    end_row = max(header_row + 1, header_row + len(rows))
    end_col = max(1, len(headers))
    ref = f"A{header_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=f"Table{ws.title.replace(' ', '')[:20]}", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    ws.freeze_panes = "C5" if end_col > 2 else "A5"
    if widths:
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
    return ws


def build_xlsx(outdir, config, index, pairs, clusters, members):
    pair_types = ["EXACT_DUPLICATE", "NEAR_DUPLICATE", "SAME_CAPTURE", "SAME_OBJECT", "RELATED_ONLY", "UNRELATED"]
    statuses = ["VALID", "NEED_REVIEW", "SPLIT_RECOMMENDED"]
    type_counts = {kind: sum(row.get("relationship_type") == kind for row in pairs) for kind in pair_types}
    status_counts = {status: sum(row.get("cluster_status") == status for row in clusters) for status in statuses}
    enriched_pairs = [{**row, "score_band": row.get("duplicate_confidence") or score_band(row, config)} for row in pairs]
    index_by_id = {row["image_id"]: row for row in index}
    canonical_code = {row["duplicate_group_id"]: index_by_id.get(row["canonical_image"], {}).get("kode_l1", "") for row in clusters}
    ultg_by_group = {}
    for group in clusters:
        gid = group["duplicate_group_id"]
        ultg_by_group[gid] = "; ".join(sorted({m["ultg"] for m in members if m["duplicate_group_id"] == gid}))
    enriched_clusters = [{**row, "canonical_code_l1": canonical_code[row["duplicate_group_id"]], "ultg": ultg_by_group[row["duplicate_group_id"]]} for row in clusters]

    wb = Workbook()
    summary = wb.active
    summary.title = "Ringkasan Lapis 3"
    summary.sheet_view.showGridLines = False
    summary["A2"] = "Relasi dan Cluster Duplikasi, Lapis 3"
    summary["A2"].font = Font(name="Arial", size=16, bold=True, color="173B58")
    summary["A3"] = "Analisis graph hubungan antar gambar. Hasil otomatis membutuhkan review pada cluster berisiko."
    summary["A3"].font = Font(name="Arial", size=10, italic=True, color="607487")
    summary["A5"] = "Indikator"; summary["B5"] = "Jumlah"
    for cell in summary[5]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B58")
    summary_rows = [
        ("Pair relationship", len(pairs)), ("Cluster", len(clusters)), ("Foto dalam cluster", len(members)),
        ("Cluster valid", status_counts["VALID"]), ("Cluster perlu review", status_counts["NEED_REVIEW"]),
        ("Cluster disarankan split", status_counts["SPLIT_RECOMMENDED"]), ("Top-K konfigurasi", config["retrieval"]["top_k"]),
        ("Model embedding", config["embedding"]["model_name"]),
    ]
    for row_no, (label, value) in enumerate(summary_rows, 6):
        summary.cell(row_no, 1, label); summary.cell(row_no, 2, value)
    summary["D5"] = "Status cluster"; summary["E5"] = "Jumlah"
    summary["G5"] = "Tipe relationship"; summary["H5"] = "Jumlah"
    for cell in (summary["D5"], summary["E5"], summary["G5"], summary["H5"]):
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="173B58")
    for row_no, status in enumerate(statuses, 6): summary.cell(row_no, 4, status); summary.cell(row_no, 5, status_counts[status])
    for row_no, kind in enumerate(pair_types, 6): summary.cell(row_no, 7, kind); summary.cell(row_no, 8, type_counts[kind])
    for row in summary.iter_rows(min_row=6, max_row=13, min_col=1, max_col=8):
        for cell in row: cell.font = Font(name="Arial", size=10, color="203042")
    summary.column_dimensions["A"].width = 30; summary.column_dimensions["B"].width = 28
    for col in ("D", "E", "G", "H"): summary.column_dimensions[col].width = 24

    pair_headers = ["image_a","image_b","embedding_similarity","phash_distance","dhash_distance","feature_match_score","relationship_score","score_band","relationship_type","duplicate_confidence","duplicate_evidence","orb_status"]
    cluster_headers = ["duplicate_group_id","canonical_image","canonical_code_l1","ultg","jumlah_gambar","average_similarity","minimum_similarity","cluster_confidence","cluster_status","cluster_validation_notes","strong_edge_ratio","cluster_diameter","relationship_type_dominant","member_images"]
    member_headers = ["duplicate_group_id","image_id","nama_file","path_rel","ultg","kode_l1","is_canonical","quality_score"]
    setup_sheet(wb.create_sheet("Pair Relationship"), "Pair Relationship", pair_headers, enriched_pairs, {"A": 48, "B": 48, "K": 58, "L": 16})
    setup_sheet(wb.create_sheet("Cluster"), "Cluster", cluster_headers, enriched_clusters, {"A": 22, "B": 48, "J": 42, "N": 90})
    setup_sheet(wb.create_sheet("Cluster Members"), "Cluster Members", member_headers, members, {"A": 22, "B": 48, "C": 38, "D": 48})
    guide_rows = [
        {"Parameter": "Urutan aman", "Nilai": "DINOv2 → HNSW Top-K → ORB → relationship table → NetworkX cluster → cluster validation", "Catatan": "Tidak memakai brute-force O(N²)."},
        {"Parameter": "Embedding model", "Nilai": config["embedding"]["model_name"], "Catatan": "Representasi visual global untuk kandidat awal."},
        {"Parameter": "Top-K", "Nilai": config["retrieval"]["top_k"], "Catatan": "Jumlah kandidat per gambar sebelum deduplikasi pasangan."},
        {"Parameter": "ORB", "Nilai": f"{config['orb']['nfeatures']} fitur, ratio {config['orb']['ratio_test']}", "Catatan": "Evidence lokal untuk scene/objek fisik yang sama."},
        {"Parameter": "Relationship threshold", "Nilai": config["relationship"]["edge_threshold"], "Catatan": "Edge graph dibuat di atas ambang ini."},
        {"Parameter": "GPS/waktu", "Nilai": "nullable", "Catatan": "Jika tidak tersedia, evidence mencatat unknown."},
        {"Parameter": "VALID", "Nilai": "Cluster dengan kemiripan internal cukup kuat", "Catatan": "Masih dapat diaudit manual."},
        {"Parameter": "NEED_REVIEW", "Nilai": "Cluster borderline atau evidence berlawanan", "Catatan": "Prioritas review manual."},
        {"Parameter": "SPLIT_RECOMMENDED", "Nilai": "Similarity chaining atau minimum similarity rendah", "Catatan": "Jangan dipakai sebagai cluster final sebelum review."},
        {"Parameter": "Batas", "Nilai": "Hipotesis otomatis", "Catatan": "Bukan keputusan inspeksi final tanpa validasi manusia."},
    ]
    setup_sheet(wb.create_sheet("Parameter dan Batas"), "Parameter dan batas interpretasi", ["Parameter", "Nilai", "Catatan"], guide_rows, {"A": 28, "B": 62, "C": 82})
    output = outdir / "Hasil_Evaluasi_Lapis_3_Relasi_Duplikasi_Foto_Isolator.xlsx"
    wb.save(output)
    return output


def build_html(outdir, config, index, pairs, clusters, members):
    pair_types = ["EXACT_DUPLICATE", "NEAR_DUPLICATE", "SAME_CAPTURE", "SAME_OBJECT", "RELATED_ONLY", "UNRELATED"]
    statuses = ["VALID", "NEED_REVIEW", "SPLIT_RECOMMENDED"]
    confidences = ["VERY_HIGH", "HIGH", "MEDIUM", "LOW"]
    index_by_id = {row["image_id"]: row for row in index}
    members_by_group = {}
    for member in members: members_by_group.setdefault(member["duplicate_group_id"], []).append(member)
    canonical_code = {row["duplicate_group_id"]: index_by_id.get(row["canonical_image"], {}).get("kode_l1", "") for row in clusters}
    ultg_by_group = {gid: "; ".join(sorted({m["ultg"] for m in group_members})) for gid, group_members in members_by_group.items()}
    cluster_data = [{**row, "canonical_code_l1": canonical_code.get(row["duplicate_group_id"], ""), "ultg": ultg_by_group.get(row["duplicate_group_id"], ""), "members": members_by_group.get(row["duplicate_group_id"], [])} for row in clusters]
    pair_data = [{**row, "score_band": row.get("duplicate_confidence") or score_band(row, config)} for row in pairs]
    type_counts = {kind: sum(row.get("relationship_type") == kind for row in pairs) for kind in pair_types}
    status_counts = {status: sum(row.get("cluster_status") == status for row in clusters) for status in statuses}
    confidence_counts = {confidence: sum(row["score_band"] == confidence for row in pair_data) for confidence in confidences}
    payload = json.dumps({"clusters": cluster_data, "pairs": pair_data, "typeCounts": type_counts, "statusCounts": status_counts, "confidenceCounts": confidence_counts}, ensure_ascii=False).replace("<", "\\u003c")
    output = outdir / "Dashboard_Evaluasi_Lapis_3_Relasi_Duplikasi_Foto_Isolator.html"
    html = f'''<!doctype html><html lang="id"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Dashboard Lapis 3</title><style>
:root{{--navy:#173b58;--blue:#176b9a;--muted:#607487;--line:#d6e3eb;--bg:#f3f7fa;--card:#fff;--ink:#203042;--amber:#a65e00;--red:#b42318;--green:#167a55}}*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:14px Arial,sans-serif}}main{{max-width:1540px;margin:auto;padding:24px}}header{{display:flex;justify-content:space-between;gap:16px;align-items:center;border-bottom:1px solid var(--line);padding-bottom:16px}}h1{{margin:0;color:var(--navy);font-size:25px}}h2{{color:var(--navy);font-size:18px}}.tag{{background:#fff1d8;color:var(--amber);border-radius:6px;padding:6px 9px;font-size:12px}}.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:18px 0}}.card,.panel{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:16px}}.card span{{display:block;color:var(--muted);font-size:12px}}.card strong{{font-size:28px;color:var(--navy)}}.filters{{display:grid;grid-template-columns:repeat(6,1fr);gap:10px;margin:18px 0}}select,input{{font:inherit;padding:9px;border:1px solid #bdcbd6;border-radius:6px;background:#fff;color:var(--navy);min-width:0}}.grid{{display:grid;grid-template-columns:.9fr 1.4fr;gap:16px}}.tablewrap{{overflow:auto;max-height:560px}}table{{width:100%;border-collapse:collapse;font-size:12px}}th{{background:var(--navy);color:#fff;padding:9px;text-align:left;position:sticky;top:0}}td{{border-bottom:1px solid #e4edf2;padding:8px;white-space:nowrap}}.pill{{border-radius:999px;padding:3px 8px;font-size:11px;font-weight:bold}}.VALID{{background:#dcfce7;color:#166534}}.NEED_REVIEW{{background:#fff1d8;color:#8a4b00}}.SPLIT_RECOMMENDED{{background:#fee2e2;color:#991b1b}}.thumbs{{display:grid;grid-template-columns:repeat(auto-fill,minmax(120px,1fr));gap:10px}}.thumb{{border:1px solid var(--line);border-radius:8px;padding:8px;background:#fff}}.thumb img{{width:100%;height:120px;object-fit:contain;background:#f7fafc}}.note{{color:var(--muted);font-size:12px;line-height:1.5}}.bar{{height:8px;background:#e5eef4;border-radius:4px;overflow:hidden}}.bar>i{{display:block;height:100%;background:var(--blue)}}.meta{{display:grid;grid-template-columns:repeat(3,1fr);gap:8px;margin:12px 0}}.meta b{{display:block;color:var(--navy);font-size:17px}}.meta span{{color:var(--muted);font-size:11px}}@media(max-width:1050px){{.filters{{grid-template-columns:repeat(3,1fr)}}}}@media(max-width:900px){{.cards,.filters,.grid,.meta{{grid-template-columns:1fr}}main{{padding:14px}}}}
</style></head><body><main><header><div><h1>Dashboard Lapis 3 Relasi Duplikasi</h1><p class="note">DINOv2 embedding → HNSW Top-K → ORB verification → relationship table → NetworkX cluster → cluster validation</p></div><div><span class="tag">Analisis otomatis · review manusia wajib</span><br><a href="Hasil_Evaluasi_Lapis_3_Relasi_Duplikasi_Foto_Isolator.xlsx">Buka Excel</a></div></header><section class="cards" id="cards"></section><section class="panel"><h2>Filter dan ringkasan evidence</h2><div class="filters"><select id="status"><option value="">Semua status cluster</option></select><select id="type"><option value="">Semua tipe relasi</option></select><select id="confidence"><option value="">Semua confidence</option></select><select id="canonicalCode"><option value="">Semua kode L1 canonical</option></select><select id="ultg"><option value="">Semua ULTG</option></select><input id="search" placeholder="Cari cluster / foto"></div><div id="bands" class="meta"></div></section><div class="grid"><section class="panel"><h2>Daftar cluster</h2><div class="tablewrap"><table><thead><tr><th>Cluster</th><th>Status</th><th>Jumlah</th><th>Avg</th><th>Min</th></tr></thead><tbody id="clusterRows"></tbody></table></div></section><section class="panel"><h2 id="detailTitle">Detail cluster</h2><p id="detailNote" class="note"></p><div id="detailMeta" class="meta"></div><div id="thumbs" class="thumbs"></div><h2>Evidence pair</h2><div class="tablewrap"><table><thead><tr><th>A</th><th>B</th><th>Score</th><th>Band</th><th>Tipe</th><th>Evidence</th></tr></thead><tbody id="pairRows"></tbody></table></div></section></div><p class="note">GPS dan waktu bersifat opsional; evidence mencatat unknown bila metadata tidak tersedia. Cluster SPLIT_RECOMMENDED hanya rekomendasi review untuk mencegah similarity chaining, bukan pemisahan otomatis.</p></main><script>
const data={payload};const {{clusters,pairs,typeCounts,statusCounts,confidenceCounts}}=data;const nf=new Intl.NumberFormat('id-ID');const pf=new Intl.NumberFormat('id-ID',{{maximumFractionDigits:3}});const esc=x=>String(x??'').replace(/[&<>"']/g,c=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[c]));const photoUrl=p=>p?'../../'+p.replaceAll(String.fromCharCode(92),'/').split('/').map(encodeURIComponent).join('/'):'#';const $=id=>document.getElementById(id);let selected=null;function opt(id,vals){{$('{id}').insertAdjacentHTML('beforeend',[...new Set(vals.filter(Boolean).flatMap(v=>String(v).split('; ')))].sort().map(v=>'<option value="'+esc(v)+'">'+esc(v)+'</option>').join(''))}}opt('status',clusters.map(c=>c.cluster_status));opt('type',pairs.map(p=>p.relationship_type));opt('confidence',pairs.map(p=>p.duplicate_confidence||p.score_band));opt('canonicalCode',clusters.map(c=>c.canonical_code_l1));opt('ultg',clusters.map(c=>c.ultg));function filtered(){{const s=$('status').value,t=$('type').value,c=$('confidence').value,k=$('canonicalCode').value,u=$('ultg').value,q=$('search').value.toLowerCase();return clusters.filter(x=>(!s||x.cluster_status===s)&&(!t||pairs.some(p=>p.relationship_type===t&&x.members.some(m=>m.image_id===p.image_a||m.image_id===p.image_b)))&&(!c||pairs.some(p=>(p.duplicate_confidence||p.score_band)===c&&x.members.some(m=>m.image_id===p.image_a||m.image_id===p.image_b)))&&(!k||x.canonical_code_l1===k)&&(!u||x.ultg.split('; ').includes(u))&&(!q||(x.duplicate_group_id+' '+x.canonical_image+' '+x.member_images+' '+x.ultg).toLowerCase().includes(q)))}}function render(){{const f=filtered();$('cards').innerHTML=[['Cluster',clusters.length],['Foto dalam cluster',clusters.reduce((a,b)=>a+Number(b.jumlah_gambar||0),0)],['Perlu review',statusCounts.NEED_REVIEW||0],['Disarankan split',statusCounts.SPLIT_RECOMMENDED||0]].map(([l,v])=>'<div class="card"><span>'+l+'</span><strong>'+nf.format(v)+'</strong></div>').join('');$('bands').innerHTML=['VERY_HIGH','HIGH','MEDIUM','LOW'].map(b=>'<div><span>'+b+' pair</span><b>'+nf.format(confidenceCounts[b]||0)+'</b><div class="bar"><i style="width:'+Math.min(100,(confidenceCounts[b]||0)/Math.max(1,pairs.length)*100)+'%"></i></div></div>').join('');$('clusterRows').innerHTML=f.map(c=>'<tr data-id="'+esc(c.duplicate_group_id)+'"><td>'+esc(c.duplicate_group_id)+'</td><td><span class="pill '+esc(c.cluster_status)+'">'+esc(c.cluster_status)+'</span></td><td>'+nf.format(c.jumlah_gambar)+'</td><td>'+pf.format(c.average_similarity)+'</td><td>'+pf.format(c.minimum_similarity)+'</td></tr>').join('');if(!selected&&f.length)selected=f[0].duplicate_group_id;if(selected&&!f.some(x=>x.duplicate_group_id===selected))selected=f[0]?.duplicate_group_id||null;renderDetail()}}function renderDetail(){{const c=clusters.find(x=>x.duplicate_group_id===selected);if(!c){{$('detailTitle').textContent='Detail cluster';$('detailNote').textContent='Tidak ada cluster sesuai filter.';$('detailMeta').innerHTML='';$('thumbs').innerHTML='';$('pairRows').innerHTML='';return}}const memberSet=new Set(c.members.map(m=>m.image_id));const cpairs=pairs.filter(p=>memberSet.has(p.image_a)&&memberSet.has(p.image_b));$('detailTitle').textContent=c.duplicate_group_id+' · '+c.cluster_status;$('detailNote').textContent='Canonical: '+c.canonical_image+' · kode L1 '+c.canonical_code_l1+' · '+c.relationship_type_dominant+' · '+c.cluster_validation_notes;$('detailMeta').innerHTML=[['Anggota',c.jumlah_gambar],['Average similarity',pf.format(c.average_similarity)],['Minimum similarity',pf.format(c.minimum_similarity)],['Strong edge ratio',pf.format(c.strong_edge_ratio)],['Cluster diameter',pf.format(c.cluster_diameter)],['ULTG',esc(c.ultg)]].map(([l,v])=>'<div><span>'+l+'</span><b>'+v+'</b></div>').join('');$('thumbs').innerHTML=c.members.map(m=>'<div class="thumb"><a href="'+photoUrl(m.path_rel)+'"><img src="'+photoUrl(m.path_rel)+'" alt="Thumbnail '+esc(m.nama_file)+'" onerror="this.alt=\'Foto tidak terbuka\'"></a><b>'+esc(m.nama_file)+'</b><p class="note">'+esc(m.kode_l1)+(m.is_canonical==='YA'?' · canonical':'')+'</p></div>').join('');$('pairRows').innerHTML=cpairs.map(p=>'<tr><td>'+esc(p.image_a)+'</td><td>'+esc(p.image_b)+'</td><td>'+pf.format(p.relationship_score)+'</td><td>'+esc(p.score_band)+'</td><td>'+esc(p.relationship_type)+'</td><td>'+esc(p.duplicate_evidence)+'</td></tr>').join('')}}for(const id of ['status','type','confidence','canonicalCode','ultg'])$(id).onchange=render;$('search').oninput=render;$('clusterRows').onclick=e=>{{const tr=e.target.closest('tr[data-id]');if(tr){{selected=tr.dataset.id;renderDetail()}}}};render();</script></body></html>'''
    output.write_text(html, encoding="utf-8")
    return output


def main():
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    work = ROOT / config["paths"]["work_dir"]
    outdir = ROOT / config["paths"]["output_dir"]
    outdir.mkdir(parents=True, exist_ok=True)
    index = read_json(work / "image_index.json")
    pairs = read_json(work / "pair_relationships.json")
    clusters = read_json(work / "clusters.json")
    members = read_json(work / "cluster_members.json")
    xlsx = build_xlsx(outdir, config, index, pairs, clusters, members)
    html = build_html(outdir, config, index, pairs, clusters, members)
    print(json.dumps({"xlsx": str(xlsx), "html": str(html), "clusters": len(clusters), "pairs": len(pairs), "images": len(index)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
