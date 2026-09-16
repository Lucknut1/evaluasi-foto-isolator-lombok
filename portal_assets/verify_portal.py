"""Read-only verification of portal metrics, downloads, and workbook data."""
from pathlib import Path
from collections import Counter
import csv
import hashlib
import json
from urllib.request import urlopen, Request
from urllib.error import HTTPError
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
BASE='http://127.0.0.1:8501'
def api(path):
    with urlopen(BASE+path,timeout=20) as r:return json.load(r)
def rows(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
s=api('/api/summary')
l1=rows(ROOT/'outputs/lapis1_work/hasil_lapis1.csv');l2=rows(ROOT/'outputs/lapis2_work/hasil_lapis2.csv')
assert s['l1']['total']==len(l1)==26299 and s['l2']['total']==len(l2)==24818
assert s['l1']['codes']==dict(Counter(r['kode'] for r in l1))
assert s['l2']['sample']==sum(r['sampel_audit_1500']=='YA' for r in l2)==1500
assert s['l3']['pairs']==368535 and s['l3']['clusters']==941 and s['l3']['members']==2644
for status,count in [('VALID',3),('NEED_REVIEW',122),('SPLIT_RECOMMENDED',816)]:
    d=api('/api/lapis3/clusters?status='+status);assert d['total']==count and all(r['cluster_status']==status for r in d['rows'])
assert api('/api/lapis3/clusters')['total']==941
assert api('/api/lapis3/clusters?type=EXACT_DUPLICATE')['total']==0
assert api('/api/lapis3/pairs')['total']==368535
assert api('/api/lapis3/pairs?type=NEAR_DUPLICATE')['total']==5876
assert api('/api/lapis3/pairs?type=UNRELATED')['total']==355016
a=api('/api/lapis3/pairs?page=0')['rows'];b=api('/api/lapis3/pairs?page=1')['rows']
assert len(a)==len(b)==50 and {r['id'] for r in a}.isdisjoint(r['id'] for r in b)
cluster=api('/api/lapis3/cluster/DUP-00001');ids={m['image_id'] for m in cluster['members']}
assert all(p['image_a'] in ids and p['image_b'] in ids for p in api('/api/lapis3/pairs?group=DUP-00001')['rows'])
for path in ['/Dataset_master.xlsx','/memory.md','/photos/%2e%2e/Dataset_master.xlsx','/outputs/portal_work/relationships.sqlite3']:
    try:urlopen(BASE+path,timeout=5);raise AssertionError('Internal file exposed: '+path)
    except HTTPError as e:assert e.code==404
file=ROOT/'outputs/portal/Ringkasan_Eksekutif_Lapis_1_2_3.xlsx'
with urlopen(BASE+'/download/excel',timeout=30) as r:
    assert 'attachment' in r.headers['Content-Disposition']
    assert hashlib.sha256(r.read()).digest()==hashlib.sha256(file.read_bytes()).digest()
for i in [1,2,3]:
    with urlopen(Request(BASE+f'/download/lapis{i}',method='HEAD'),timeout=10) as r:assert r.status==200 and int(r.headers['Content-Length'])>1000
w=openpyxl.load_workbook(file,read_only=True,data_only=True)
assert w.sheetnames==['Eksekutif Summary','Lapis 1','Lapis 2','Lapis 3']
summary=w.worksheets[0]
for cell,value in [('C6',26299),('C10',24818),('C15',368535),('C16',941),('C17',2644)]:assert summary[cell].value==value
assert abs(summary['H18'].value-3/941)<1e-9
detail1=list(w['Lapis 1'].iter_rows(min_row=5,values_only=True));detail2=list(w['Lapis 2'].iter_rows(min_row=5,values_only=True));detail3=list(w['Lapis 3'].iter_rows(min_row=5,values_only=True))
assert len(detail1)==26299 and len(detail2)==24818 and len(detail3)==2644
assert {(r[0],r[1]):r[2] for r in detail1}=={(r['ultg'],r['nama_file']):r['kode'] for r in l1}
assert sum(r[3]=='YA' for r in detail2)==1500
assert len({r[0] for r in detail3})==941
assert len({(r[0],r[10],r[11]) for r in detail3})==2644
for sheet in w:
    assert not any(c.data_type=='e' for row in sheet.iter_rows() for c in row)
w.close()
guard=json.loads((ROOT/'outputs/lapis2_work/guardrail_internal.json').read_text())
assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==sha for p,sha in guard['l1_sha256'].items())
print('PASS: source totals, four sheets, detail rows, unchanged L1, filters, pagination, download bytes, and internal-route restrictions.')
